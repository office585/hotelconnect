import os
import csv
import json
import time
import tempfile
from datetime import datetime
from pathlib import Path

from playwright.sync_api import sync_playwright
from openpyxl import load_workbook

from google.oauth2.service_account import Credentials
from googleapiclient.discovery import build
from googleapiclient.http import MediaFileUpload


# ============================================================
# ALAP BEÁLLÍTÁSOK
# ============================================================

LOGIN_URL = "https://p.hotelconnect.hu/login"

EMAIL = os.environ["HOTELCONNECT_EMAIL"]
PASSWORD = os.environ["HOTELCONNECT_PASSWORD"]

GDRIVE_SERVICE_ACCOUNT_JSON = os.environ[
    "GDRIVE_SERVICE_ACCOUNT_JSON"
]

WAIT_COMPANY = 3
WAIT_NAV = 2

CURRENT_YEAR = datetime.now().year


# ============================================================
# HOTELCONNECT CÉG -> GOOGLE DRIVE MAPPA
# ============================================================

COMPANIES = {
    "Maverick Athenaeum": {
        "code": "ATH",
        "folder_id": "19bo5GiU6lrPEgfrSbJrO74e7pWD9Ng7U",
    },

    "Maverick Downtown": {
        "code": "DT",
        "folder_id": "1HjE1CMPEIYHqG6HA7aPc5OqG0GAfHdXU",
    },

    "Maverick City Lodge": {
        "code": "SOHO",
        "folder_id": "1fz1PwvGgmam-vZpg9SF3chn7s4ujTUYf",
    },

    "Giselle Vintage Doubles": {
        "code": "GVD",
        "folder_id": "1WWJ3dhu1yw2Lfw3ZxqTqaRtLsjLmg1ac",
    },

    "Maverick Urban Lodge": {
        "code": "CENTRAL",
        "folder_id": "1OlJCdki0z-TC1lrewUrL4f0nPzGAbpwO",
    },

    "Giselle Buda Castle": {
        "code": "GBC",
        "folder_id": "1xk3SOqhKWNPbRMkgZM1JleGwgYHN8p2K",
    },

    "The Amberlyn Suite Hotel": {
        "code": "AMBERLYN",
        "folder_id": "102qpagWkmb8j9NO7IU93D6qTVVYdBGKt",
    },
}


# ============================================================
# GOOGLE DRIVE KAPCSOLAT
# ============================================================

def get_drive_service():

    service_account_info = json.loads(
        GDRIVE_SERVICE_ACCOUNT_JSON
    )

    credentials = Credentials.from_service_account_info(
        service_account_info,
        scopes=[
            "https://www.googleapis.com/auth/drive"
        ],
    )

    drive_service = build(
        "drive",
        "v3",
        credentials=credentials,
        cache_discovery=False,
    )

    return drive_service


# ============================================================
# GOOGLE DRIVE FELTÖLTÉS
# ============================================================

def upload_or_replace_file(
    drive_service,
    local_file,
    folder_id,
    drive_filename,
):

    escaped_name = drive_filename.replace(
        "'",
        "\\'"
    )

    query = (
        f"name = '{escaped_name}' "
        f"and '{folder_id}' in parents "
        f"and trashed = false"
    )

    result = (
        drive_service.files()
        .list(
            q=query,
            fields="files(id,name)",
            supportsAllDrives=True,
            includeItemsFromAllDrives=True,
        )
        .execute()
    )

    existing_files = result.get(
        "files",
        []
    )

    media = MediaFileUpload(
        str(local_file),
        mimetype="text/csv",
        resumable=True,
    )

    if existing_files:

        file_id = existing_files[0]["id"]

        print(
            f"Drive fájl frissítése: "
            f"{drive_filename}"
        )

        (
            drive_service.files()
            .update(
                fileId=file_id,
                media_body=media,
                supportsAllDrives=True,
            )
            .execute()
        )

    else:

        print(
            f"Új Drive fájl létrehozása: "
            f"{drive_filename}"
        )

        metadata = {
            "name": drive_filename,
            "parents": [folder_id],
        }

        (
            drive_service.files()
            .create(
                body=metadata,
                media_body=media,
                fields="id",
                supportsAllDrives=True,
            )
            .execute()
        )


# ============================================================
# XLSX -> NYERS CSV
# ============================================================

def xlsx_to_csv(
    xlsx_path,
    csv_path
):

    print(
        "XLSX átalakítása nyers CSV-vé..."
    )

    workbook = load_workbook(
        filename=xlsx_path,
        read_only=True,
        data_only=True,
    )

    worksheet = workbook.active

    with open(
        csv_path,
        "w",
        newline="",
        encoding="utf-8-sig",
    ) as csv_file:

        writer = csv.writer(
            csv_file,
            delimiter=",",
            quoting=csv.QUOTE_MINIMAL,
        )

        for row in worksheet.iter_rows(
            values_only=True
        ):

            cleaned_row = []

            for value in row:

                if value is None:
                    cleaned_row.append("")
                else:
                    cleaned_row.append(value)

            writer.writerow(
                cleaned_row
            )

    workbook.close()


# ============================================================
# FŐ PROGRAM
# ============================================================

def run():

    drive_service = get_drive_service()

    temp_dir = Path(
        tempfile.mkdtemp()
    )

    print(
        "Ideiglenes könyvtár:",
        temp_dir
    )

    successful = []
    failed = []

    with sync_playwright() as p:

        # FONTOS:
        # ugyanaz a mód, mint a lokálisan működő verzióban
        browser = p.chromium.launch(
            headless=False
        )

        context = browser.new_context(
            viewport={
                "width": 1440,
                "height": 900,
            },
            accept_downloads=True,
        )

        page = context.new_page()

        page.set_default_timeout(
            60000
        )

        # ====================================================
        # LOGIN
        # ====================================================

        print(
            "Login oldal megnyitása..."
        )

        page.goto(
            LOGIN_URL,
            wait_until="domcontentloaded"
        )

        page.locator(
            'input[type="email"]'
        ).fill(
            EMAIL
        )

        page.locator(
            'input[type="password"]'
        ).fill(
            PASSWORD
        )

        print(
            "Bejelentkezés..."
        )

        page.locator(
            'button[type="submit"]'
        ).click()

        # ====================================================
        # LOGIN UTÁNI OLDAL MEGVÁRÁSA
        #
        # NEM URL-T NÉZÜNK,
        # HANEM AZ EGYÉRTELMŰEN BELÉPÉS UTÁNI ELEMET.
        # ====================================================

        page.locator(
            'aside button[title*="NTAK jelentés"]'
        ).first.wait_for(
            state="visible",
            timeout=60000
        )

        print(
            "Sikeres belépés."
        )

        print(
            "Aktuális URL:",
            page.url
        )

        print(
            f"Várakozás login után: "
            f"{WAIT_COMPANY} mp"
        )

        time.sleep(
            WAIT_COMPANY
        )

        # ====================================================
        # HOTELCONNECT EGYSÉGEK KIOLVASÁSA
        # ====================================================

        company_buttons = page.locator(
            'aside button[title*="NTAK jelentés"]'
        )

        company_count = (
            company_buttons.count()
        )

        print(
            f"\nTalált HotelConnect egységek: "
            f"{company_count}"
        )

        found_companies = []

        for i in range(
            company_count
        ):

            button = (
                company_buttons.nth(i)
            )

            company_name = (
                button
                .locator("span.truncate")
                .inner_text()
                .strip()
            )

            found_companies.append(
                company_name
            )

            print(
                f"  {i + 1}. "
                f"{company_name}"
            )

        # ====================================================
        # CSAK A KÉRT 7 CÉGEN MEGYÜNK VÉGIG
        # ====================================================

        total_companies = len(
            COMPANIES
        )

        for index, (
            company_name,
            config
        ) in enumerate(
            COMPANIES.items(),
            start=1
        ):

            print(
                "\n" + "=" * 65
            )

            print(
                f"[{index}/{total_companies}] "
                f"{company_name}"
            )

            print(
                "=" * 65
            )

            if (
                company_name
                not in found_companies
            ):

                print(
                    "HIBA: ez az egység nincs "
                    "a HotelConnect listában."
                )

                failed.append(
                    company_name
                )

                continue

            try:

                # ============================================
                # CÉG KIVÁLASZTÁSA
                # ============================================

                print(
                    f"Cég kiválasztása: "
                    f"{company_name}"
                )

                company_button = (
                    page
                    .locator(
                        'aside button'
                    )
                    .filter(
                        has=page.locator(
                            "span.truncate",
                            has_text=company_name
                        )
                    )
                )

                company_button.click()

                print(
                    f"Várakozás cégváltás után: "
                    f"{WAIT_COMPANY} mp"
                )

                time.sleep(
                    WAIT_COMPANY
                )

                # ============================================
                # XLSX LETÖLTÉS MENÜ
                # ============================================

                print(
                    "XLSX letöltése menü megnyitása..."
                )

                xlsx_buttons = (
                    page.get_by_role(
                        "button",
                        name="XLSX letöltése",
                        exact=True
                    )
                )

                xlsx_buttons.first.click()

                time.sleep(
                    WAIT_NAV
                )

                # ============================================
                # EZ AZ ÉV
                # ============================================

                print(
                    "'Ez az év' kiválasztása..."
                )

                page.get_by_role(
                    "button",
                    name="Ez az év",
                    exact=True
                ).click()

                time.sleep(
                    WAIT_NAV
                )

                # ============================================
                # MODÁLIS XLSX LETÖLTÉS
                # ============================================

                print(
                    "XLSX export indítása..."
                )

                modal_xlsx_button = (
                    page.get_by_role(
                        "button",
                        name="XLSX letöltése",
                        exact=True
                    ).last
                )

                with page.expect_download(
                    timeout=180000
                ) as download_info:

                    modal_xlsx_button.click()

                download = (
                    download_info.value
                )

                # ============================================
                # IDEIGLENES XLSX
                # ============================================

                xlsx_filename = (
                    f"{config['code']}_"
                    f"{CURRENT_YEAR}.xlsx"
                )

                xlsx_path = (
                    temp_dir
                    / xlsx_filename
                )

                download.save_as(
                    str(xlsx_path)
                )

                xlsx_size_mb = (
                    xlsx_path.stat().st_size
                    / 1024
                    / 1024
                )

                print(
                    f"XLSX letöltve: "
                    f"{xlsx_size_mb:.2f} MB"
                )

                # ============================================
                # CSV KÉSZÍTÉSE
                # ============================================

                csv_filename = (
                    f"hotelconnect_"
                    f"{config['code']}_"
                    f"{CURRENT_YEAR}.csv"
                )

                csv_path = (
                    temp_dir
                    / csv_filename
                )

                xlsx_to_csv(
                    xlsx_path,
                    csv_path
                )

                csv_size_mb = (
                    csv_path.stat().st_size
                    / 1024
                    / 1024
                )

                print(
                    f"CSV elkészült: "
                    f"{csv_size_mb:.2f} MB"
                )

                # ============================================
                # IDEIGLENES XLSX TÖRLÉSE
                # ============================================

                try:

                    xlsx_path.unlink()

                    print(
                        "Ideiglenes XLSX törölve."
                    )

                except Exception as e:

                    print(
                        "XLSX törlési figyelmeztetés:",
                        e
                    )

                # ============================================
                # DRIVE FELTÖLTÉS
                # ============================================

                print(
                    "Feltöltés Google Drive-ra..."
                )

                upload_or_replace_file(
                    drive_service=drive_service,
                    local_file=csv_path,
                    folder_id=config[
                        "folder_id"
                    ],
                    drive_filename=csv_filename,
                )

                print(
                    f"KÉSZ: {company_name}"
                )

                successful.append(
                    company_name
                )

                # ============================================
                # IDEIGLENES CSV TÖRLÉSE
                # ============================================

                try:

                    csv_path.unlink()

                except Exception:
                    pass

                time.sleep(
                    WAIT_NAV
                )

            except Exception as e:

                print(
                    f"HIBA ennél a cégnél: "
                    f"{company_name}"
                )

                print(
                    str(e)
                )

                failed.append(
                    company_name
                )

                try:

                    cancel_button = (
                        page.get_by_role(
                            "button",
                            name="Mégse",
                            exact=True
                        )
                    )

                    if (
                        cancel_button.is_visible()
                    ):

                        cancel_button.click()

                        time.sleep(
                            WAIT_NAV
                        )

                except Exception:
                    pass

                continue

        browser.close()

    # ========================================================
    # ÖSSZEGZÉS
    # ========================================================

    print(
        "\n" + "=" * 65
    )

    print(
        "FUTÁS VÉGE"
    )

    print(
        "=" * 65
    )

    print(
        "\nSikeres cégek:"
    )

    for company in successful:

        print(
            f"OK: {company}"
        )

    if failed:

        print(
            "\nHibás cégek:"
        )

        for company in failed:

            print(
                f"HIBA: {company}"
            )

        raise RuntimeError(
            f"{len(failed)} cég "
            f"feldolgozása sikertelen."
        )

    print(
        "\nMinden cég sikeresen elkészült."
    )


if __name__ == "__main__":
    run()
